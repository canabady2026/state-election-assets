from openpyxl import load_workbook, Workbook
import glob
import os

# Get all xlsx files
files = sorted(glob.glob("*.xlsx"))

# Create output workbook
out_wb = Workbook()

# Remove default empty sheet
out_wb.remove(out_wb.active)

for file in files:
    wb = load_workbook(file, data_only=True)

    # Get first sheet
    ws = wb.worksheets[0]

    # Create sheet in output workbook
    new_ws = out_wb.create_sheet(title=os.path.splitext(os.path.basename(file))[0][:31])

    # Copy values
    for row in ws.iter_rows(values_only=True):
        new_ws.append(row)

    print(f"Added first sheet from {file}")

# Save merged workbook
out_wb.save("merged.xlsx")

print("Saved merged.xlsx")