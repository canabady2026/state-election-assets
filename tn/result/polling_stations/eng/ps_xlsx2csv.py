#!/usr/bin/env python3

import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

INPUT_XLSX = "acs-ps.xlsx"
OUTPUT_DIR = "csv_output"

Path(OUTPUT_DIR).mkdir(exist_ok=True)

wb = load_workbook(INPUT_XLSX, data_only=True)

for sheet_name in wb.sheetnames:

    ws = wb[sheet_name]

    records = []

    # AC001 -> 1
    m = re.search(r"(\d+)", sheet_name)
    ac_code = int(m.group(1)) if m else None

    for row in ws.iter_rows(values_only=True):

        if not row:
            continue

        # Replace None
        row = [x if x is not None else "" for x in row]

        # Skip empty rows
        if not any(str(x).strip() for x in row):
            continue

        first_col = str(row[0]).strip().lower()

        # Skip headers
        if (
            "sl" in first_col
            or "polling station" in first_col
            or first_col == ""
        ):
            continue

        # Skip numbering rows like 1 2 3 4
        if row[:4] == [1, 2, 3, 4]:
            continue

        # First column must be numeric
        if not isinstance(row[0], (int, float)):
            continue

        sl_no = int(row[0])

        # ---------------------------------------------------
        # FORMAT 1
        # sl_no | ps_number | polling_station | areas | type
        # ---------------------------------------------------

        if len(row) >= 5 and row[4] != "":

            polling_station_number = row[1]

            if polling_station_number in ["", None]:
                polling_station_number = sl_no

            polling_station = row[2]
            polling_areas = row[3]
            polling_station_type = row[4]

        # ---------------------------------------------------
        # FORMAT 2 (AC028 type)
        # sl_no | polling_station | areas | type
        # ---------------------------------------------------

        else:

            polling_station_number = sl_no
            polling_station = row[1] if len(row) > 1 else ""
            polling_areas = row[2] if len(row) > 2 else ""
            polling_station_type = row[3] if len(row) > 3 else ""

        records.append({
            "sl_no": sl_no,
            "ac_code": ac_code,
            "polling_station_number": polling_station_number,
            "polling_station": polling_station,
            "polling_areas": polling_areas,
            "polling_station_type": polling_station_type,
        })

    df = pd.DataFrame(records, columns=[
        "sl_no",
        "ac_code",
        "polling_station_number",
        "polling_station",
        "polling_areas",
        "polling_station_type",
    ])

    out_file = Path(OUTPUT_DIR) / f"{sheet_name.lower()}.csv"

    df.to_csv(out_file, index=False)

    print(f"Created: {out_file}")

print("\nDone.")